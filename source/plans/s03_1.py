# -*- coding: utf-8 -*-
from common import *
from schedule import ROWS, NOTE
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def short(link):
    if not link: return ''
    parts = [p for p in link.split('/') if p]
    return 'ikea.bg/…/' + parts[-1]


def price_txt(p):
    if p is None: return '—'
    if isinstance(p, str): return p
    return fmt(p)


def sheet():
    import textwrap
    S = Sheet('03.1', 'Мебели – спецификация', scale='—')
    ax = S.ax
    X0, top = 22.0, 285.0
    cols = [('№', 7), ('Помещение', 15), ('Елемент', 104), ('Продукт / изпълнение', 66),
            ('Материал / цвят', 94), ('Ш × Д × В, cm', 27), ('Бр.', 7), ('Ед. цена, €', 16),
            ('Сума, €', 15), ('Връзка', 37)]
    W = sum(c[1] for c in cols)
    FS = 4.5
    cw = FS * 0.61 * PT2MM          # approx. average glyph width, mm
    LH, PAD = 2.4, 0.8
    hh = 7
    def cells(r):
        key, room, item, prod, fin, dims, q, pr, link = r
        return [str(key) if key != '' else '', room, item, prod, fin, dims, str(q),
                price_txt(pr), fmt(pr * q) if isinstance(pr, float) else '', short(link)]
    wrapped = []
    for r in ROWS:
        cs = cells(r)
        ws = [textwrap.wrap(v, max(4, int((cols[j][1] - 2.5) / cw))) or [''] if j in (2, 3, 4) else [v] for j, v in enumerate(cs)]
        wrapped.append(ws)
    hs = [PAD + LH * max(len(w) for w in ws) for ws in wrapped]
    H = hh + sum(hs) + PAD + LH
    ax.add_patch(Rectangle((X0, top - H), W, H, fc='white', ec='k', lw=0.6, zorder=15))
    ax.add_patch(Rectangle((X0, top - hh), W, hh, fc='0.88', ec='k', lw=0.5, zorder=15.2))
    x = X0
    xs = []
    for (t, w) in cols:
        xs.append(x)
        S.ptext(x + 1.2, top - hh / 2, t, 5.2, weight='bold', va='center', zorder=17)
        if x > X0:
            ax.plot([x, x], [top - H, top], 'k', lw=0.25, zorder=16)
        x += w
    total = 0.0
    prev_room = None
    y = top - hh
    for i, (r, ws, h) in enumerate(zip(ROWS, wrapped, hs)):
        room = r[1]; pr = r[7]; q = r[6]
        if room != prev_room and i:
            ax.plot([X0, X0 + W], [y, y], 'k', lw=0.45, zorder=16)
        elif i % 2:
            ax.add_patch(Rectangle((X0 + 0.2, y - h), W - 0.4, h, fc='0.95', ec='none', zorder=15.1))
        if isinstance(pr, float): total += pr * q
        for j, lines in enumerate(ws):
            v = lines if room != prev_room or j != 1 else ['']
            ha = 'right' if j in (7, 8) else ('center' if j in (0, 6) else 'left')
            xx = xs[j] + (cols[j][1] - 1.2 if ha == 'right' else (cols[j][1] / 2 if ha == 'center' else 1.2))
            if len(v) == 1:
                S.ptext(xx, y - PAD / 2 - LH / 2, v[0], FS, va='center', ha=ha, zorder=17,
                        weight='bold' if j in (0, 1) else 'normal', color='#1F4E79' if j == 9 else 'k')
            else:
                S.ptext(xx, y - PAD / 2 - 0.15, '\n'.join(v), FS, va='top', ha=ha, zorder=17, color='k',
                        linespacing=1.12)
        y -= h
        prev_room = room
    ax.plot([X0, X0 + W], [y, y], 'k', lw=0.6, zorder=16)
    yt = y - (PAD + LH) / 2
    S.ptext(xs[2] + 1.2, yt, 'ОБЩО (само позиции с цена от ikea.bg)', 5.0, weight='bold', va='center', zorder=17)
    S.ptext(xs[8] + cols[8][1] - 1.2, yt, fmt(total), 5.0, weight='bold', va='center', ha='right', zorder=17)
    ny = top - H - 3.5
    S.ptext(X0, ny, 'Бележки: ' + NOTE.replace('; „по поръчка“', ';\n„по поръчка“'), 4.6, va='top', zorder=17, linespacing=1.4)
    S.ptext(X0, ny - 8.5, 'Пълните връзки са в „Мебели - спецификация.xlsx“. Чертежи на мебелите по поръчка – листове 03.2 – 03.4.\n'
            'Номерата (№) съответстват на лист 02 Мебелировка.', 4.6, va='top', zorder=17, linespacing=1.4)
    print('table bottom', round(top - H, 1))
    S.title_block()
    S.save('03.1 Мебели - спецификация')
    return total


def xlsx():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Спецификация'
    ws['A1'] = 'Апартамент 75, Младост 4 – Мебели, спецификация'; ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f'Дата: {DATE}.  ' + NOTE; ws['A2'].font = Font(italic=True, size=9)
    hdr = ['№ (лист 02)', 'Помещение', 'Елемент', 'Продукт / изпълнение', 'Материал / цвят', 'Ш × Д × В, cm',
           'Бр.', 'Ед. цена, €', 'Сума, €', 'Връзка']
    thin = Side(style='thin', color='999999')
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDDDDD')
        c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = Border(bottom=thin, top=thin)
    r0 = 5
    for i, r in enumerate(ROWS):
        key, room, item, prod, fin, dims, q, pr, link = r
        rr = r0 + i
        vals = [key if key != '' else None, room, item, prod, fin, dims, q,
                pr if pr is not None else ('—' if pr is None else pr), None, link or None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=rr, column=j, value=v)
            c.alignment = Alignment(wrap_text=j in (3, 4, 5), vertical='center')
            c.border = Border(bottom=thin)
        if isinstance(pr, float):
            ws.cell(row=rr, column=9, value=f'=G{rr}*H{rr}')
            ws.cell(row=rr, column=8).number_format = '#,##0.00'
            ws.cell(row=rr, column=9).number_format = '#,##0.00'
        if link:
            ws.cell(row=rr, column=10).hyperlink = link
            ws.cell(row=rr, column=10).font = Font(color='1F4E79', underline='single')
    rl = r0 + len(ROWS)
    ws.cell(row=rl, column=3, value='ОБЩО (позиции с цена)').font = Font(bold=True)
    c = ws.cell(row=rl, column=9, value=f'=SUM(I{r0}:I{rl - 1})'); c.font = Font(bold=True); c.number_format = '#,##0.00'
    for j, w in enumerate([10, 11, 44, 30, 50, 18, 6, 12, 12, 70], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f'A4:J{rl - 1}'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0
    p = os.path.join(OUT, 'Мебели - спецификация.xlsx'); wb.save(p)


if __name__ == '__main__':
    t = sheet(); xlsx(); print(round(t, 2))
